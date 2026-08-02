#!/usr/bin/env python3
"""Build the Town of Glover redesign summary as a polished Word guide."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_FILES = Path(__file__).resolve().parents[1]
ROOT = PROJECT_FILES / "townofglover-content-archive"
SOURCE = ROOT / "REDESIGN-CONTENT-SUMMARY.md"
OUTPUT = ROOT / "REDESIGN-CONTENT-SUMMARY.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(32, 55, 72)
GRAY = RGBColor(90, 98, 108)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, url, label=None):
    label = label or url
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.extend([r_fonts, color, underline])
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_numbering(doc):
    numbering = doc.part.numbering_part.element
    used_abstract = {
        int(x.get(qn("w:abstractNumId")))
        for x in numbering.findall(qn("w:abstractNum"))
    }
    used_num = {
        int(x.get(qn("w:numId")))
        for x in numbering.findall(qn("w:num"))
    }
    abstract_bullet = max(used_abstract | {0}) + 1
    abstract_decimal = abstract_bullet + 1
    num_bullet = max(used_num | {0}) + 1
    num_decimal = num_bullet + 1

    def add_abstract(abstract_id, fmt, text, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        level.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        level.append(lvl_text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "tab")
        level.append(suffix)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        level.append(lvl_jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        level.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            level.append(r_pr)
        abstract.append(level)
        numbering.append(abstract)

    def add_num(num_id, abstract_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    add_abstract(abstract_bullet, "bullet", "•", "Calibri")
    add_abstract(abstract_decimal, "decimal", "%1.")
    add_num(num_bullet, abstract_bullet)
    add_num(num_decimal, abstract_decimal)

    for style_name, num_id in (("List Bullet", num_bullet), ("List Number", num_decimal)):
        style = doc.styles[style_name]
        p_pr = style._element.get_or_add_pPr()
        old = p_pr.find(qn("w:numPr"))
        if old is not None:
            p_pr.remove(old)
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_node])
        p_pr.append(num_pr)


def add_rich_text(paragraph, text):
    # Render Markdown bold spans and URLs while keeping the parser deliberately
    # simple and deterministic.
    pattern = re.compile(r"(\*\*.+?\*\*|https?://\S+|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("http"):
            add_hyperlink(paragraph, token.rstrip(".,);"), token.rstrip(".,);"))
            trailing = token[len(token.rstrip(".,);")):]
            if trailing:
                set_run_font(paragraph.add_run(trailing))
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color=GRAY)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]))


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Lead Callout" not in doc.styles:
        style = doc.styles.add_style("Lead Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Lead Callout"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = NAVY
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(10)
    style.paragraph_format.left_indent = Inches(0.18)
    style.paragraph_format.right_indent = Inches(0.18)
    style.paragraph_format.line_spacing = 1.2

    if "Source Note" not in doc.styles:
        style = doc.styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Source Note"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)
    style.font.color.rgb = GRAY
    style.paragraph_format.space_before = Pt(2)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.1


def style_callout(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "2E74B5")
    borders.append(left)
    p_pr.append(borders)


def add_cover(doc):
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("WEBSITE CONTENT GUIDE"), size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("Town of Glover"), size=30, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    set_run_font(
        p.add_run("Reorganized Content Summary for Website Redesign"),
        size=15,
        color=DARK_BLUE,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shading)
    metrics = [
        ("733", "public pages"),
        ("1,782", "files and images"),
        ("296,644", "extracted words"),
    ]
    for idx, (value, label) in enumerate(metrics):
        if idx:
            set_run_font(p.add_run("     |     "), size=10, color=GRAY)
        set_run_font(p.add_run(value), size=15, color=NAVY, bold=True)
        set_run_font(p.add_run(f" {label}"), size=9.5, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("Source captured July 30, 2026"), size=11, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        p.add_run("A planning guide; the complete original archive remains unchanged."),
        size=9.5,
        color=GRAY,
        italic=True,
    )
    p.add_run().add_break(WD_BREAK.PAGE)


def add_contents_page(doc):
    doc.add_heading("Contents", level=1)
    entries = [
        "1. Essential Town Information",
        "2. Common Resident Tasks",
        "3. Town Government",
        "4. Meetings and Public Records",
        "5. Services and Public Safety",
        "6. Community, Recreation, and Local Organizations",
        "7. Lakes, Rivers, Flooding, and Environmental Projects",
        "8. News, Announcements, and Historical Content",
        "9. Recommended Structure for the Redesigned Website",
        "10. Content Cleanup and Migration Priorities",
        "11. Suggested Migration Sequence",
        "12. Archive Reference",
    ]
    for entry in entries:
        p = doc.add_paragraph(style="List Number")
        # Numbering is already part of the displayed source text, so use a plain
        # paragraph styled for matching indent without creating duplicate markers.
        p.style = doc.styles["Normal"]
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(7)
        set_run_font(p.add_run(entry), color=NAVY, bold=True)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    configure_numbering(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("TOWN OF GLOVER  |  WEBSITE CONTENT GUIDE"), size=8.5, color=GRAY, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("Page "), size=8.5, color=GRAY)
    add_field(footer, "PAGE")

    add_cover(doc)
    add_contents_page(doc)

    in_quote = False
    first_title_seen = False
    for raw in lines:
        line = raw.rstrip()
        if not line:
            in_quote = False
            continue
        if line == "---":
            continue
        if line.startswith("# "):
            if not first_title_seen:
                first_title_seen = True
                continue
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            title = line[3:]
            if title in {"A reorganized guide for website redesign", "How to use this document"}:
                if title == "How to use this document":
                    doc.add_heading(title, level=1)
                continue
            doc.add_heading(title, level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("> "):
            p = doc.add_paragraph(style="Lead Callout")
            add_rich_text(p, line[2:])
            style_callout(p)
            in_quote = True
        elif re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
        elif line.startswith("**Source") or line.startswith("**Primary source"):
            p = doc.add_paragraph(style="Source Note")
            add_rich_text(p, line)
        else:
            p = doc.add_paragraph()
            add_rich_text(p, line)

    doc.core_properties.title = "Town of Glover Website Content Summary"
    doc.core_properties.subject = "Reorganized content guide for website redesign"
    doc.core_properties.author = "Town of Glover website content archive"
    doc.core_properties.keywords = "Glover, Vermont, website redesign, content inventory"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
